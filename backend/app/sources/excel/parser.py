"""Adapters for MISIS institute grids and a normalized row-based XLSX table.

Never execute spreadsheet macros/formulas. Keep original text and provenance.
Unsupported layouts fail the import instead of replacing good data with emptiness.
"""
import hashlib
import io
import json
import re
import zipfile
from dataclasses import asdict, dataclass
from typing import Optional

import openpyxl
import xlrd

DAY_NAMES = ['понедельник', 'вторник', 'среда', 'четверг', 'пятница', 'суббота', 'воскресенье']
EN_DAYS = ['monday', 'tuesday', 'wednesday', 'thursday', 'friday', 'saturday', 'sunday']
GROUP = re.compile(r'(?<![А-ЯЁA-Z])[БМСА][А-ЯЁA-Z]{1,12}-\d{2}(?:-[А-ЯЁA-Z0-9]+)*', re.I)
TIME = re.compile(r'(?<!\d)([0-2]?\d)[.:]([0-5]\d)(?::[0-5]\d)?\s*[-–—]\s*([0-2]?\d)[.:]([0-5]\d)(?::[0-5]\d)?')
PERSON = re.compile(r'[А-ЯЁA-Z][а-яёa-zА-ЯЁA-Z-]+\s+[А-ЯЁA-Z]\s*\.\s*[А-ЯЁA-Z]\s*\.')
TYPE = re.compile(r'\((Лекционные|Практические|Лабораторные|Лекция|Практика|Семинар|Лаб\.?|Лек\.?|Пр\.?)\)', re.I)
SUBGROUP = re.compile(r'(\d(?:\s*[,и]\s*\d)*)\s*(?:п\.?\s*г\.?|подгрупп[аы]?)', re.I)


class ParseError(ValueError):
    pass


def clean(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace('\xa0', ' ').strip()


def day_number(value):
    normalized = re.sub(r'[^а-яёa-z]', '', clean(value).lower())
    for names in (DAY_NAMES, EN_DAYS):
        candidate = normalized.translate(str.maketrans({'а': 'a', 'е': 'e', 'о': 'o', 'с': 'c', 'у': 'y', 'т': 't'})) if names is EN_DAYS else normalized
        for i, name in enumerate(names):
            if candidate == name or candidate[::-1] == name:
                return i
    return None


def time_range(value):
    match = TIME.search(clean(value))
    if not match:
        return None
    h1, m1, h2, m2 = map(int, match.groups())
    if h1 > 23 or h2 > 23 or (h1, m1) >= (h2, m2):
        raise ParseError(f'Invalid time: {value}')
    return f'{h1:02}:{m1:02}', f'{h2:02}:{m2:02}'


@dataclass
class Lesson:
    group_name: str
    weekday: int
    start_time: str
    end_time: str
    subject: str
    lesson_type: str = ''
    teacher: str = ''
    room: str = ''
    subgroup: str = ''
    week_pattern: str = 'all'
    notes: str = ''
    raw_text: str = ''
    source_url: str = ''
    source_sheet: str = ''
    source_cell: str = ''

    def record(self):
        value = asdict(self)
        identity = {k: v for k, v in value.items() if not k.startswith('source_') and k != 'raw_text'}
        value['id'] = hashlib.sha256(json.dumps(identity, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
        return value


@dataclass
class Sheet:
    name: str
    cells: list
    spans: dict

    def value(self, row, col):
        if row < 0 or col < 0 or row >= len(self.cells) or col >= len(self.cells[row]):
            return ''
        top, _, left, _ = self.spans.get((row, col), (row, row + 1, col, col + 1))
        return clean(self.cells[top][left])

    def span(self, row, col):
        return self.spans.get((row, col), (row, row + 1, col, col + 1))


def load_sheets(data):
    if data.startswith(b'PK'):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 150 * 1024 * 1024:
                raise ParseError('Expanded workbook exceeds 150 MB')
        book = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        raw = [(s.title, [[c.value for c in row] for row in s],
                [(m.min_row - 1, m.max_row, m.min_col - 1, m.max_col) for m in s.merged_cells.ranges]) for s in book]
        book.close()
    elif data.startswith(b'\xd0\xcf\x11\xe0'):
        book = xlrd.open_workbook(file_contents=data, formatting_info=True)
        raw = [(s.name, [s.row_values(i) for i in range(s.nrows)], s.merged_cells) for s in book.sheets()]
        book.release_resources()
    else:
        raise ParseError('Response is not an XLSX or XLS workbook')
    for name, cells, merged in raw:
        if sum(len(row) for row in cells) > 2_000_000:
            raise ParseError(f'Sheet too large: {name}')
        spans = {}
        for span in merged:
            top, bottom, left, right = span
            for row in range(top, bottom):
                for col in range(left, right):
                    spans[(row, col)] = span
        yield Sheet(name, cells, spans)


def content(raw):
    lines = [line.strip() for line in raw.splitlines() if line.strip()]
    first = lines[0] if lines else ''
    kind = TYPE.search(first)
    teachers = []
    notes = []
    # A full initials-form name is extracted even when it shares the subject line.
    for line in lines:
        teachers.extend(m.group(0) for m in PERSON.finditer(line))
    subject = PERSON.sub('', first)
    subject = TYPE.sub('', subject)
    sub = SUBGROUP.search(subject)
    subject = SUBGROUP.sub('', subject).strip(' ,;')
    for line in lines[1:]:
        remainder = PERSON.sub('', line).strip(' ,;')
        if remainder:
            notes.append(remainder)
    return dict(subject=subject, teacher=', '.join(dict.fromkeys(teachers)),
                lesson_type=kind.group(1) if kind else '', subgroup=sub.group(0) if sub else '',
                notes='\n'.join(notes))


def flat_table(sheet, source):
    aliases = {'группа': 'group_name', 'день': 'weekday', 'день недели': 'weekday', 'начало': 'start_time', 'конец': 'end_time', 'предмет': 'subject', 'дисциплина': 'subject', 'тип': 'lesson_type', 'преподаватель': 'teacher', 'аудитория': 'room', 'подгруппа': 'subgroup', 'неделя': 'week_pattern', 'примечание': 'notes'}
    for row_idx, row in enumerate(sheet.cells[:15]):
        header = {aliases.get(clean(v).lower(), clean(v).lower()): i for i, v in enumerate(row) if clean(v)}
        required = {'group_name', 'weekday', 'start_time', 'end_time', 'subject'}
        if not required <= header.keys():
            continue
        lessons = []
        groups = set()
        for r in range(row_idx + 1, len(sheet.cells)):
            if not any(clean(v) for v in sheet.cells[r]):
                continue
            values = {key: sheet.value(r, col) for key, col in header.items() if key in Lesson.__dataclass_fields__}
            if not all(values.get(k) for k in required):
                raise ParseError(f'{sheet.name}, row {r + 1}: incomplete lesson')
            day = day_number(values['weekday'])
            if day is None and values['weekday'] in [str(i) for i in range(7)]:
                day = int(values['weekday'])
            times = time_range(f"{values['start_time']}-{values['end_time']}")
            if day is None or times is None:
                raise ParseError(f'{sheet.name}, row {r + 1}: invalid day/time')
            values.update(weekday=day, start_time=times[0], end_time=times[1])
            week = values.get('week_pattern', 'all') or 'all'
            week = {'нечётная': 'odd', 'нечетная': 'odd', 'чётная': 'even', 'четная': 'even', 'все': 'all'}.get(week.lower(), week.lower())
            if week not in ('all', 'odd', 'even'):
                raise ParseError(f'{sheet.name}, row {r + 1}: invalid week')
            values.update(week_pattern=week, source_url=source, source_sheet=sheet.name, source_cell=f'A{r + 1}', raw_text=' | '.join(sheet.value(r, c) for c in header.values()))
            lessons.append(Lesson(**values))
            groups.add(values['group_name'])
        return lessons, groups
    return None


def grid_table(sheet, source, upper_row_week):
    header_row = None
    groups = []
    for r, row in enumerate(sheet.cells[:20]):
        found = [(c, GROUP.search(clean(v)).group(0).upper()) for c, v in enumerate(row) if GROUP.search(clean(v))]
        if found:
            header_row, groups = r, found
            break
    if header_row is None:
        return None
    # The time column is detected from actual values, not fixed Excel coordinates.
    counts = {c: sum(bool(time_range(sheet.value(r, c))) for r in range(header_row + 1, min(len(sheet.cells), 100))) for c in range(groups[0][0])}
    if not counts or max(counts.values()) == 0:
        raise ParseError(f'{sheet.name}: group headers found but no supported time column')
    time_col = max(counts, key=counts.get)
    active_rows = [r for r in range(header_row + 1, len(sheet.cells)) if time_range(sheet.value(r, time_col))]
    blocks = []
    for r in active_rows:
        t = time_range(sheet.value(r, time_col))
        direct_day = day_number(sheet.value(r, 0))
        if not blocks or t[0] < time_range(sheet.value(blocks[-1][-1], time_col))[0] or (direct_day is not None and direct_day != day_number(sheet.value(blocks[-1][0], 0)) and sheet.span(r, 0)[0] == r):
            blocks.append([])
        blocks[-1].append(r)
    lessons = []
    for block in blocks:
        labels = [day_number(sheet.value(r, 0)) for r in block]
        day = next((d for d in labels if d is not None), None)
        if day is None:
            # Some IBO sheets write the weekday bottom-to-top, one letter per row.
            day = day_number(''.join(clean(sheet.cells[r][0]) for r in block))
        if day is None:
            if any(sheet.value(r, c) for r in block for c, _ in groups):
                raise ParseError(f'{sheet.name}, row {block[0] + 1}: weekday cannot be identified')
            continue
        for gi, (col, group) in enumerate(groups):
            next_col = groups[gi + 1][0] if gi + 1 < len(groups) else sheet.span(header_row, col)[3]
            width = next_col - col
            if width <= 1:
                width = groups[1][0] - groups[0][0] if len(groups) > 1 else 2
            if width < 2 or width > 12 or width % 2:
                raise ParseError(f'{sheet.name}: unsupported group width {width} for {group}')
            for lane in range(col, col + width, 2):
                previous: Optional[Lesson] = None
                for r in block:
                    raw = sheet.value(r, lane)
                    if not raw or raw.strip() in ('-', '—'):
                        previous = None
                        continue
                    top, bottom, _, _ = sheet.span(r, lane)
                    if top < r and top in block:
                        continue
                    times = time_range(sheet.value(r, time_col))
                    end = time_range(sheet.value(min(bottom - 1, block[-1]), time_col))
                    fields = content(raw)
                    room = sheet.value(r, lane + 1)
                    if sheet.span(r, lane + 1) == sheet.span(r, lane):
                        room = ''
                    if TIME.fullmatch(raw) and previous:
                        explicit_times = time_range(raw)
                        previous.start_time, previous.end_time = explicit_times
                        previous.raw_text += '\n' + raw
                        continue
                    if not fields['subject'] and fields['teacher'] and previous:
                        # Teacher-only second half of a 45 + 45 minute IBO pair.
                        previous.teacher = ', '.join(dict.fromkeys(filter(None, [previous.teacher, fields['teacher']])))
                        previous.end_time = max(previous.end_time, times[1])
                        previous.room = room or previous.room
                        previous.raw_text += '\n' + raw + ('\n' + room if room else '')
                        continue
                    if not fields['subject']:
                        raise ParseError(f'{sheet.name}, row {r + 1}: teacher without subject')
                    time_top, time_bottom, _, _ = sheet.span(r, time_col)
                    week = 'all'
                    if time_bottom - time_top == 2 and bottom - top < 2:
                        week = upper_row_week if r == time_top else ('even' if upper_row_week == 'odd' else 'odd')
                    if width > 2 and not fields['subgroup']:
                        fields['subgroup'] = f'{(lane - col) // 2 + 1} подгруппа'
                    previous = Lesson(group_name=group, weekday=day, start_time=times[0], end_time=(end or times)[1],
                                      room=room, week_pattern=week, raw_text=raw + ('\n' + room if room else ''),
                                      source_url=source, source_sheet=sheet.name, source_cell=f'{openpyxl.utils.get_column_letter(lane + 1)}{r + 1}', **fields)
                    lessons.append(previous)
    return lessons, {g for _, g in groups}


def consolidate(lessons):
    # Identical numerator/denominator records describe one weekly lesson.
    buckets = {}
    for lesson in lessons:
        key = tuple((k, v) for k, v in asdict(lesson).items() if k not in ('week_pattern', 'source_cell'))
        if key in buckets and buckets[key].week_pattern != lesson.week_pattern:
            buckets[key].week_pattern = 'all'
        else:
            buckets[key] = lesson
    return list({item.record()['id']: item for item in buckets.values()}.values())


def parse_workbook(data, source_url='', upper_row_week='odd'):
    lessons, groups, warnings = [], set(), []
    for sheet in load_sheets(data):
        result = flat_table(sheet, source_url)
        if result is None:
            result = grid_table(sheet, source_url, upper_row_week)
        if result is None:
            if any(clean(v) for row in sheet.cells for v in row):
                warnings.append(f'Справочный лист: {sheet.name}')
            continue
        sheet_lessons, sheet_groups = result
        lessons.extend(sheet_lessons)
        groups.update(sheet_groups)
    if not lessons or not groups:
        raise ParseError('Workbook has no recognizable lessons; previous data preserved')
    return consolidate(lessons), groups, warnings
